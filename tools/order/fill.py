# 기Log가 할 일을 그대로 흉내 내요: 자리 표시 글자를 바꾸고, 이름이 "위치도"·"현장사진"인 도형 자리에 그림을 맞춰 넣어요.
import sys, json
from pptx import Presentation
from PIL import Image, ImageOps
tpl, out, data_json = sys.argv[1:4]
d = json.loads(data_json)
p = Presentation(tpl)
s = p.slides[0]
def frames(shapes):
    for sh in shapes:
        if sh.has_text_frame: yield sh.text_frame
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells: yield c.text_frame
missing = set('{%s}' % k for k in d['text'])
for tf in frames(s.shapes):
    for para in tf.paragraphs:
        full = ''.join(r.text for r in para.runs)
        if '{' not in full: continue
        for k, v in d['text'].items():
            if '{%s}' % k in full:
                missing.discard('{%s}' % k); full = full.replace('{%s}' % k, v)
        lines = full.split('\n')
        para.runs[0].text = lines[0]
        for r in para.runs[1:]: r._r.getparent().remove(r._r)
        # 줄바꿈은 같은 서식의 문단을 더해요
        import copy
        prev = para._p
        for ln in lines[1:]:
            np_ = copy.deepcopy(para._p); prev.addnext(np_); prev = np_
            from pptx.text.text import _Paragraph
            q = _Paragraph(np_, para._parent); q.runs[0].text = ln
for sh in list(s.shapes):
    if sh.name in d['images']:
        x, y, w, h = sh.left, sh.top, sh.width, sh.height
        sh._element.getparent().remove(sh._element)
        img = ImageOps.exif_transpose(Image.open(d['images'][sh.name])).convert('RGB')
        tmp = f'/tmp/claude-0/wo2/_{sh.name}.jpg'; img.save(tmp, quality=90)
        iw, ih = img.size; sc = min(w / iw, h / ih)
        W, H = int(iw * sc), int(ih * sc)
        s.shapes.add_picture(tmp, x + (w - W) // 2, y, W, H)
p.save(out)
print('없는 자리:', sorted(missing) or '없음')
